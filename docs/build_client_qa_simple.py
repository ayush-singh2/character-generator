"""Build a simple question-and-answer Word document (.docx).

Plain format: a short title, then each client question in bold followed by its
answer in normal text. No cards, tables, or brand styling.

    ../.venv/bin/python docs/build_client_qa_simple.py
"""
import os

from docx import Document
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.realpath(__file__))
OUT = os.path.join(HERE, "Client-Answers-Simple-QA.docx")

QA = [
    ("How much choice does the author have over what's on each page? Will the tool "
     "just draw based on the manuscript, or will the author be able to provide "
     "specific instructions?",
     [
         "The author has full choice. The tool does not just draw whatever the manuscript "
         "says — it treats the manuscript as the story and the author's instructions as the "
         "direction. You can be as hands-off or as hands-on as you like, page by page.",
         "If you write nothing extra, the tool reads each scene and produces a complete, "
         "coherent book on its own. But for any page you can also add specific instructions in "
         "your own words — for example, “wide shot, Ella kneeling beside the crate, warm "
         "evening light, cat in the background.” The tool treats those notes as "
         "authoritative and builds the page around them.",
         "So you decide what is on each page, who is in it, how it is framed, and the mood — "
         "either by letting the story speak for itself or by giving as much detail as you want.",
     ]),

    ("Style and font choices: Will the author pick their desired illustration style from a "
     "set of samples? Same question for the font, size, and placement of the text.",
     [
         "Yes to both. The illustration style, the character look, and all text settings are "
         "author choices made from clear samples and previews — not decisions the tool makes "
         "for you.",
         "Illustration style: you pick the overall style from a set of samples (for example soft "
         "watercolour, bold flat colour, painterly storybook, line-and-wash). Once chosen, it is "
         "applied consistently across the whole book.",
         "Character look: before the book is built, the tool creates a reference sheet for each "
         "main character — face, outfit, colours, and signature items. You approve or adjust that "
         "look first, and only then does the tool illustrate the pages using your approved "
         "character. This is what keeps a character looking like the same character on every page.",
         "Fonts and text: the font is chosen from a curated, print-safe set; the size is a "
         "comfortable default you can adjust; and the text is placed into a calm, uncluttered "
         "area of each illustration so it stays easy to read.",
     ]),

    ("Corrections and updates: How will illustration edits work after spreads have already "
     "been created? Say the author wants the character to have a hat on, or the cat to be "
     "orange, or to add a jump rope to a scene. Will they be able to make those updates and "
     "keep everything else as is?",
     [
         "Yes — this is a core feature. After a page has been illustrated, you can request a "
         "specific change with a simple instruction, and the tool applies only that change while "
         "keeping everything else exactly as it was.",
         "Your examples are precisely the kinds of edits it is built for:",
         "•  “Add a red hat to the character” — the hat is added; the face, pose, "
         "outfit, and the whole scene stay the same.",
         "•  “Make the cat orange” — only the cat's colour changes; everything "
         "else is untouched.",
         "•  “Add a jump rope to the scene” — the prop is introduced into the "
         "existing scene without redrawing the page from scratch.",
         "Corrections are applied on top of the existing illustration, so the parts you already "
         "approved are preserved. Each change affects only that one page, and you can keep "
         "refining a single page until it is exactly right. If a change should carry through the "
         "book (say the character now always wears the hat), it can be locked into the "
         "character so every page stays consistent.",
     ]),

    ("Any limitations on the final tool? How many books can we create with this a month? "
     "How much upkeep will be required?",
     [
         "There is no fixed monthly cap built into the tool. How many books you can make depends "
         "on two practical things: your own review time (the book moves as fast as you approve "
         "characters and pages) and the per-image cost of the image service. A single author "
         "working steadily can comfortably produce several books a month, and more if pages are "
         "approved in batches. The tool can also work on multiple books at once.",
         "Upkeep for you as the author is minimal — you supply a manuscript, make your choices, "
         "and review pages. There is no server to run and no software to maintain between books. "
         "Behind the scenes, ongoing maintenance (keeping the image and text services current, "
         "adding new style samples or fonts, and refinements as the models improve) is handled "
         "on the development side.",
         "Honest limitations: the AI is directed, not perfect on the first try, so a page will "
         "occasionally need one correction pass — which is exactly what the per-page edit "
         "feature is for. It is fair to leave a little room for error here: not everything will "
         "land perfectly the first time, and there will naturally be some small imperfections "
         "along the way. Text positioning in particular cannot always be placed exactly where "
         "you might picture it — the tool aims for a clean, readable, uncluttered spot in the "
         "art, but it can sit slightly off, and a minor nudge during the review pass usually "
         "settles it. These are small, manageable issues rather than blockers, and the per-page "
         "edit loop is there to smooth them out. Nothing is off-limits; some very specific "
         "requests simply take one extra round.",
     ]),
]


def main():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    r = title.add_run("Answers to Your Questions")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)
    title.paragraph_format.space_after = Pt(4)

    sub = doc.add_paragraph()
    rs = sub.add_run("Illustrated-Book Studio  ·  18 August 2026")
    rs.font.size = Pt(10)
    rs.font.color.rgb = RGBColor(0x7C, 0x8A, 0x8C)
    sub.paragraph_format.space_after = Pt(14)

    for i, (q, answers) in enumerate(QA, 1):
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(10)
        qp.paragraph_format.space_after = Pt(5)
        rq = qp.add_run(f"Q{i}.  {q}")
        rq.bold = True
        rq.font.size = Pt(11.5)
        rq.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)

        for a in answers:
            ap = doc.add_paragraph()
            ap.paragraph_format.space_after = Pt(6)
            run = ap.add_run(a)
            run.font.size = Pt(11)

    doc.save(OUT)
    print(f"-> {OUT}")


if __name__ == "__main__":
    main()
